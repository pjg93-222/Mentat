"""contracts/flow.json v1: GET /api/flow, the Sankey of downstream spending (lane F, boards #44-#46).

Asserted against the producer (flow_data.committee_payload and the live route), on a
synthetic snapshot whose every figure is worked by hand below, and on the saved FAIRSHAKE
and Senate Leadership Fund snapshots when this machine has them. The contract's negative
controls must each be refused by the same check the producer passes (DE-021).

Scenario (cycle 2024; every amount a whole number of dollars so the arithmetic is visible):
  ROOT FUND (C90000001) receives ACME CORP 1,000,000 cash + 500,000 of bitcoin in kind (booked
  back out, never sold, so it is still held), DOE, JANE 300,000 + 200,000 of USD coin in kind
  (booked out and sold: the 200,000 of sale proceeds are not new money), a 10,000 vendor refund
  (line 15, nets against operating), 50,000 back from BRIDGE PAC (a loop), 1,000 unitemized,
  and starts with 100,000 of cash. It sends BRIDGE PAC 800,000, spends 400,000 for SMITH and
  150,000 against JONES, 60,000 on operating, and ends with 251,000.
  Counted once: 2,051,000 in + 100,000 prior cash = 800,000 + 550,000 + 50,000 operating
  + 751,000 not yet spent (251,000 cash + 500,000 bitcoin).
  BRIDGE PAC (C90000002) receives the 800,000 and 100,000 from ROE, RICK; returns 50,000 to ROOT
  FUND (a returned-to stub), passes 30,000 to ONWARD PAC (not collected), gives SMITH 20,000 and
  spends 600,000 for SMITH, 10,000 on operating, and ends with 190,000. Every payment it makes is
  dated after ROOT FUND's money arrived, so the scenario reads the same before and after F12
  (the dated walk); a separate test moves one payment earlier and accepts either regime.
Downstream committees' figures on the saved snapshots are taken from the measurement
(downstream.snapshot_view), never written down here, so F12 moving them does not break a test.
"""
from collections import OrderedDict, defaultdict
import copy
from decimal import Decimal
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import urllib.error
import urllib.request

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'app'))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
import downstream  # noqa: E402
import flow_data  # noqa: E402
import pac_flow_access  # noqa: E402

CONTRACT = (Path(__file__).resolve().parents[2] / 'audits' / 'mentat_failings_20260915'
            / 'contracts' / 'flow.json')
SAVED = Path(__file__).resolve().parents[1] / '.cache' / 'pac-evidence'
FAIRSHAKE = 'snapshot-b26761887a1f44d58cc38fdb8b720a78'
SLF = 'snapshot-28d0bc70c5bb44b799b0a4cbb3522aee'
CRYPTO = Path(r'C:\Users\pgoff\OneDrive\Desktop\Mentat Complete Files\Crypto Industry Spending Overview.xlsx')
ROOT_ID, BRIDGE_ID, ONWARD_ID = 'C90000001', 'C90000002', 'C90000003'


def receipt(sub_id, name, amount, day, line='11AI', entity='ORG', text='', contributor_id=None):
    return {'sub_id': sub_id, 'contributor_name': name, 'contributor_id': contributor_id, 'entity_type': entity,
            'contribution_receipt_amount': amount, 'contribution_receipt_date': day, 'line_number': line,
            'filing_form': 'F3X', 'receipt_type_full': text, 'memo_code': None, 'memo_text': None,
            'memoed_subtotal': False, 'pdf_url': f'https://docquery.fec.gov/cgi-bin/fecimg/?2024{sub_id[1:]:0>10}'}


def payment(sub_id, recipient, amount, day, line='21B', kind=None, committee=None, candidate=None, text=''):
    return {'sub_id': sub_id, 'recipient_name': recipient, 'recipient_committee_id': committee,
            'disbursement_amount': amount, 'disbursement_date': day, 'line_number': line, 'disbursement_type': kind,
            'candidate_id': candidate, 'candidate_name': 'SMITH, ANN' if candidate else None,
            'disbursement_description': text, 'memo_code': None, 'memoed_subtotal': False}


def spend(sub_id, candidate, name, party, direction, amount, day):
    return {'sub_id': sub_id, 'candidate_id': candidate, 'candidate_name': name, 'candidate_party': party,
            'support_oppose_indicator': direction, 'expenditure_amount': amount, 'expenditure_date': day,
            'is_notice': False, 'memo_code': None, 'memoed_subtotal': False}


def totals(name, receipts_, disbursements, cash_begin, cash_end, **extra):
    return {'committee_name': name, 'receipts': receipts_, 'disbursements': disbursements,
            'cash_on_hand_beginning_period': cash_begin, 'last_cash_on_hand_end_period': cash_end,
            'coverage_start_date': '2023-01-01T00:00:00', 'coverage_end_date': '2024-12-31T00:00:00',
            'committee_type_full': 'Super PAC (Independent Expenditure-Only)', **extra}


def fixture_units():
    root = {'committee_id': ROOT_ID, 'cycle': 2024,
            'receipts': [receipt('r1', 'ACME CORP', 1000000.0, '2023-03-01'),
                         receipt('r2', 'ACME CORP', 500000.0, '2023-04-01', text='IN-KIND - 10 BITCOIN - NOT LIQUIDATED'),
                         receipt('r3', 'DOE, JANE', 300000.0, '2023-03-02', entity='IND'),
                         receipt('r4', 'DOE, JANE', 200000.0, '2023-05-01', entity='IND',
                                 text='IN-KIND - 200,000 US DOLLAR COIN (USDC)'),
                         receipt('r5', 'COIN EXCHANGE', 200000.0, '2023-05-02', line='17',
                                 text='200,000 USDC SOLD VIA COIN EXCHANGE - PURCHASER UNKNOWN'),
                         receipt('r6', 'VENDOR LLC', 10000.0, '2023-06-01', line='15', text='REFUND OF MEDIA BUYS'),
                         receipt('r7', 'BRIDGE PAC', 50000.0, '2024-02-01', line='11C', entity='PAC',
                                 contributor_id=BRIDGE_ID, text='CONTRIBUTION')],
            'disbursements': [payment('b1', 'BRIDGE PAC', 800000.0, '2023-06-01', line='22', kind='24G', committee=BRIDGE_ID),
                              payment('b2', 'ACME CORP', 500000.0, '2023-04-01', text='IN-KIND 10 BITCOIN NOT LIQUIDATED'),
                              payment('b3', 'DOE, JANE', 200000.0, '2023-05-01', text='IN-KIND 200,000 USDC'),
                              payment('b4', 'MEDIA BUYERS', 60000.0, '2023-07-01', text='MEDIA PRODUCTION')],
            'expenditures': [spend('e1', 'H4XX01001', 'SMITH, ANN', 'REP', 'S', 400000.0, '2023-08-01'),
                             spend('e2', 'S4YY00002', 'JONES, BOB REP.', 'DEM', 'O', 150000.0, '2023-09-01')],
            'totals': totals('ROOT FUND', 2261000.0, 2110000.0, 100000.0, 251000.0,
                             offsets_to_operating_expenditures=10000.0, contribution_refunds=0.0,
                             operating_expenditures=760000.0, independent_expenditures=550000.0,
                             transfers_to_affiliated_committee=800000.0, fed_candidate_committee_contributions=0.0)}
    bridge = {'committee_id': BRIDGE_ID, 'cycle': 2024,
              'receipts': [receipt('s1', 'ROOT FUND', 800000.0, '2023-06-02', line='12', entity='PAC', contributor_id=ROOT_ID),
                           receipt('s2', 'ROE, RICK', 100000.0, '2023-06-10', entity='IND')],
              'disbursements': [payment('c1', 'ROOT FUND', 50000.0, '2024-02-01', line='23', kind='24K', committee=ROOT_ID),
                                payment('c2', 'ONWARD PAC', 30000.0, '2023-07-15', line='23', kind='24K', committee=ONWARD_ID),
                                payment('c3', 'SMITH FOR CONGRESS', 20000.0, '2023-10-01', line='23', kind='24K',
                                        committee='C90000009', candidate='H4XX01001'),
                                payment('c4', 'OFFICE RENT', 10000.0, '2023-07-01')],
              'expenditures': [spend('f1', 'H4XX01001', 'SMITH, ANN', 'REP', 'S', 600000.0, '2023-09-15')],
              'totals': totals('BRIDGE PAC', 900000.0, 710000.0, 0.0, 190000.0, offsets_to_operating_expenditures=0.0,
                               contribution_refunds=0.0, operating_expenditures=10000.0,
                               independent_expenditures=600000.0, fed_candidate_committee_contributions=100000.0)}
    return [root, bridge]


def chargeback_units():
    """The scenario with four donors' charge-backs (lane I, SLF 2022's shape: a 'CHARGED BACK' row of
    -$500 is Robert Weirman's only row in the window). DYKE and WEIRMAN net below zero (their gifts came
    before the window): $5,500 left ROOT FUND as refunds. REYNOLDS nets to zero. PART nets to $2,000.
    The summary nets them: receipts and cash on hand are $3,500 lower than the base scenario's."""
    units = fixture_units()
    root = units[0]

    def back(sub_id, name, amount, day):
        row = receipt(sub_id, name, amount, day, entity='IND', text='CONTRIBUTION')
        row['memo_text'] = 'CHARGED BACK' if amount < 0 else None
        return row
    root['receipts'] += [back('r8', 'WEIRMAN, ROB', -500.0, '2023-01-06'),
                         back('r9', 'DYKE, MERRITT', -5000.0, '2023-01-06'),
                         back('r10', 'REYNOLDS, PETER', 5000.0, '2023-09-28'),
                         back('r11', 'REYNOLDS, PETER', -5000.0, '2023-10-12'),
                         back('r12', 'PART, PAT', 3000.0, '2023-03-05'),
                         back('r13', 'PART, PAT', -1000.0, '2023-04-05')]
    root['totals']['receipts'] -= 3500.0
    root['totals']['last_cash_on_hand_end_period'] -= 3500.0
    return units


def write_fixture(directory, units=None, identifier='flow-fixture'):
    folder = Path(directory) / identifier
    folder.mkdir(parents=True, exist_ok=True)
    (folder / 'snapshots.json').write_text(json.dumps(units or fixture_units()), encoding='ascii')
    (folder / 'evidence.json').write_text(json.dumps({
        'version': 1, 'company': 'Acme', 'retrieved_at': '2026-09-26T01:10:08+00:00',
        'downstream_coverage': [{'committee_id': ONWARD_ID, 'cycle': 2024, 'status': 'collection_limit'}],
        'transfers': [{'destination': ONWARD_ID, 'destination_name': 'ONWARD PAC'}],
        'committee_finances': []}), encoding='ascii')
    return identifier


def produce(directory, identifier='flow-fixture', **options):
    return flow_data.committee_payload(ROOT_ID, '2024', identifier=identifier, root=Path(directory), **options)


def refusal(directory, **options):
    try:
        produce(directory, **options)
    except flow_data.FlowRefused as error:
        return error.body()
    raise AssertionError('the flow was not refused')


def tampered_view(directory):
    view = downstream.snapshot_view('flow-fixture', ROOT_ID, '2024', Path(directory))
    view = copy.deepcopy(view)
    view['relationships'][0]['amount'] += 1.0      # one dollar the flow does not draw
    return view


def by_id(body):
    return {node['id']: node for node in body['nodes']}


def link(body, source, target, kind=None):
    found = [item for item in body['links'] if item['source'] == source and item['target'] == target
             and (kind is None or item['kind'] == kind)]
    assert len(found) == 1, (source, target, kind, found)
    return found[0]


def contract():
    return json.loads(CONTRACT.read_text(encoding='utf-8'))


class FlowContract(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        write_fixture(cls.temp.name)
        cls.body = produce(cls.temp.name)
        cls.contract = contract()

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_the_producer_answer_holds_the_contract_and_is_its_example(self):
        self.assertEqual(flow_data.contract_problems(self.body), [])
        for example in self.contract['examples']:
            if example.get('producer') != 'fixture':
                continue
            request = example['request']
            if example['status'] == 409:
                options = {'basis': request.get('basis', 'reported')}
                if example['name'] == 'refused: the flow disagrees with the evidence network':
                    options['evidence_view'] = tampered_view(self.temp.name)
                body = refusal(self.temp.name, **options)
            else:
                body = produce(self.temp.name, top=int(request.get('top', 8)), depth=int(request.get('depth', 3)))
            self.assertEqual(flow_data.contract_problems(body, example['status']), [], example['name'])
            self.assertEqual(json.loads(json.dumps(body)), example['body'], example['name'])

    def test_the_root_counts_every_dollar_once(self):
        nodes, totals = by_id(self.body), self.body['totals']
        root = nodes['committee:' + ROOT_ID]
        self.assertEqual((root['value_in'], root['value_out']), (2151000.0, 2151000.0))
        self.assertEqual(totals['in'], 2151000.0)
        self.assertEqual(totals['counted_in'], 2051000.0)       # receipts less proceeds less the refund
        self.assertEqual(totals['not_yet_spent'], {'cash': 251000.0, 'held': 500000.0, 'total': 751000.0,
                                                   'held_kinds': ['bitcoin']})
        self.assertEqual(totals['operating_net'], 50000.0)
        self.assertEqual(totals['in_kind'], {'booked': 700000.0, 'proceeds': 200000.0, 'held': 500000.0,
                                             'excess_proceeds': 0.0})
        self.assertEqual(link(self.body, 'donor:ACME CORP', root['id'])['value'], 1500000.0)
        self.assertEqual(nodes['donor:ACME CORP']['held'], 500000.0)
        self.assertNotIn('donor:COIN EXCHANGE', nodes)           # sale proceeds are not a donor
        self.assertNotIn('donor:VENDOR LLC', nodes)              # a refund nets against operating
        self.assertEqual(link(self.body, root['id'], 'balance:operating:2')['value'], 50000.0)
        self.assertEqual(link(self.body, root['id'], 'balance:unspent:2')['value'], 751000.0)
        self.assertEqual(link(self.body, 'balance:prior_cash:0', root['id'])['value'], 100000.0)

    def test_a_loop_is_a_stub_at_both_ends_and_nothing_counts_twice(self):
        nodes = by_id(self.body)
        back = nodes[f'stub:return_in:committee:{ROOT_ID}:{BRIDGE_ID}']
        self.assertEqual((back['label'], back['value_out']), ('back from Bridge PAC', 50000.0))
        returned = nodes[f'stub:return:committee:{BRIDGE_ID}:{ROOT_ID}']
        self.assertEqual((returned['label'], returned['value_in']), ('returned to Root Fund', 50000.0))
        self.assertNotIn('donor:BRIDGE PAC', nodes)               # money from the network is not a donor
        bridge = nodes['committee:' + BRIDGE_ID]
        self.assertEqual((bridge['value_in'], bridge['value_out']), (900000.0, 900000.0))
        self.assertEqual(nodes[f'stub:external:committee:{BRIDGE_ID}:{BRIDGE_ID}']['value_out'], 100000.0)
        self.assertEqual(bridge['share_from_upstream'][0]['share'], 0.888889)

    def test_an_onward_payment_is_a_stub_with_the_measurement_s_date_flags(self):
        onward = link(self.body, 'committee:' + BRIDGE_ID, f'stub:onward:committee:{BRIDGE_ID}:{ONWARD_ID}')
        self.assertEqual(onward['value'], 30000.0)
        self.assertTrue(all('date_order_ok' not in item for item in self.body['links']))
        self.assertEqual((self.body['totals']['before_gate'], self.body['totals']['date_unverified']), (0.0, 0.0))
        self.assertEqual([row['id'] for row in self.body['coverage']['unreached']], [ONWARD_ID])
        self.assertEqual(self.body['coverage']['unreached'][0]['reason'], 'collection_limit')

    def test_money_paid_before_the_network_s_arrived_is_disclosed_not_drawn(self):
        # F12: BRIDGE PAC's payment to ONWARD PAC predates ROOT FUND's money (gate 2023-06-01), so
        # the measurement leaves it out; the flow discloses it beside BRIDGE PAC, never as a flow.
        units = fixture_units()
        early = next(row for row in units[1]['disbursements'] if row['sub_id'] == 'c2')
        early['disbursement_date'] = '2023-02-01'
        with tempfile.TemporaryDirectory() as directory:
            write_fixture(directory, units)
            body = produce(directory)
            walk = downstream.snapshot_view('flow-fixture', ROOT_ID, '2024', Path(directory))
        self.assertEqual(flow_data.contract_problems(body), [])
        self.assertFalse(any(edge['source'] == BRIDGE_ID and edge['destination'] == ONWARD_ID for edge in walk['edges']))
        self.assertTrue(any(edge['source'] == BRIDGE_ID and edge['destination'] == ONWARD_ID
                            for edge in walk['date_rule']['excluded_before_gate']))
        nodes = by_id(body)
        self.assertNotIn(f'stub:onward:committee:{BRIDGE_ID}:{ONWARD_ID}', nodes)
        gated = nodes[f'stub:before_gate:committee:{BRIDGE_ID}:{BRIDGE_ID}']
        self.assertEqual((gated['label'], gated['value_in'], gated['kind']),
                         ("Not counted: paid before the network's money arrived", 30000.0, 'balance'))
        self.assertEqual(link(body, 'committee:' + BRIDGE_ID, gated['id'])['kind'], 'uncounted')
        self.assertEqual(body['totals']['before_gate'], 30000.0)
        self.assertIn("Not counted: paid before the network's money arrived $30K",
                      [note['text'] for note in body['notes']])
        self.assertFalse(any(nodes[item['target']]['kind'] == 'candidate' and item['value'] == 30000.0
                             for item in body['links']))
        bridge = nodes['committee:' + BRIDGE_ID]
        self.assertEqual(bridge['value_in'], bridge['value_out'])

    def test_a_not_in_summary_stub_on_the_root_is_disclosed_not_counted_twice(self):
        # Spending the summary cannot carry is drawn as a 'Not in the summary' stub into its node so
        # the node balances (FAIRSHAKE 2026's root has one): the root-inflow rule sets it aside.
        # Teeth: the same inflow from any other stub is money counted twice.
        root = 'committee:' + ROOT_ID

        def with_stub(balance):
            body = copy.deepcopy(self.body)
            stub_id = f'stub:external:{root}:{ROOT_ID}'
            body['nodes'].append({'id': stub_id, 'stage': by_id(body)[root]['stage'], 'label': 'Not in the summary',
                                  'kind': 'balance', 'stub': 'external', 'attach': root, 'balance': balance,
                                  'value_in': 0.0, 'value_out': 25000.0})
            body['links'].append({'source': stub_id, 'target': root, 'value': 25000.0, 'kind': 'external_in', 'rows': 0})
            operating = link(body, root, 'balance:operating:2')
            operating['value'] = round(operating['value'] + 25000.0, 2)
            nodes = by_id(body)
            for node, key in ((nodes[root], 'value_in'), (nodes[root], 'value_out'),
                              (nodes['balance:operating:2'], 'value_in')):
                node[key] = round(node[key] + 25000.0, 2)
            return body
        self.assertEqual(flow_data.contract_problems(with_stub('unreconciled')), [])
        self.assertIn('the root takes in more or less than totals.in (money counted twice or lost)',
                      flow_data.contract_problems(with_stub('prior_cash')))

    def test_grouping_keeps_the_money_and_names_the_rest(self):
        grouped = produce(self.temp.name, top=3)
        self.assertEqual(flow_data.contract_problems(grouped), [])
        other = by_id(grouped)['other:0']
        self.assertEqual((other['label'], other['value_out'], other['member_count']), ('All other contributors', 501000.0, 2))
        self.assertEqual(sum(item['value'] for item in grouped['links'] if item['target'] == 'committee:' + ROOT_ID),
                         sum(item['value'] for item in self.body['links'] if item['target'] == 'committee:' + ROOT_ID))

    def test_attributed_without_an_estimate_bundle_is_refused(self):
        body = refusal(self.temp.name, basis='attributed')
        import graph_data
        self.assertEqual(body['error'], graph_data.NO_ESTIMATE_BUNDLE)
        self.assertEqual((body['reason'], body['offer']), ('no_estimate_bundle', 'reported'))

    def test_a_flow_that_disagrees_with_the_evidence_network_is_refused(self):
        body = refusal(self.temp.name, evidence_view=tampered_view(self.temp.name))
        self.assertEqual(body['reason'], 'reconciliation')
        self.assertFalse(body['reconciliation']['equal'])
        self.assertEqual(body['reconciliation']['candidates'], {'evidence_network': 1170001.0, 'flow': 1170000.0})

    def test_every_negative_control_is_refused(self):
        for control in self.contract['negative_controls']:
            problems = flow_data.contract_problems(control['payload'], control.get('status', 200))
            self.assertTrue(problems, control['why'])
            self.assertTrue(any(control['expect'] in problem for problem in problems), (control['why'], problems))

    def test_the_request_is_checked(self):
        for query in ({'cycle': ['2023'], 'root': ['committee:C00835959']}, {'cycle': ['2024'], 'root': ['C00835959']},
                      {'cycle': ['2024'], 'root': ['committee:C00835959'], 'top': ['99']},
                      {'cycle': ['2024'], 'root': ['committee:C00835959'], 'basis': ['traced']},
                      {'cycle': ['2024'], 'root': ['committee:C00835959'], 'id': ['../x']}):
            with self.assertRaises(ValueError):
                flow_data.parse_request(query)
        self.assertEqual(flow_data.parse_request({'cycle': ['2023-2024'], 'root': ['committee:C00835959']})['cycle'], '2024')
        with self.assertRaises(ValueError):
            flow_data.parse_request({'cycle': ['2024'], 'root': ['industry'], 'direct': ['1']})
        with self.assertRaises(ValueError):
            flow_data.parse_request({'cycle': ['2024'], 'root': ['committee:C00835959'],
                                     'direct': ['1'], 'id': ['saved']})


class DirectCommitteeFlow(unittest.TestCase):
    def test_existing_unit_is_used_without_a_snapshot_choice_or_fec_call(self):
        with tempfile.TemporaryDirectory() as directory:
            write_fixture(directory)
            with patch.object(pac_flow_access, 'direct_evidence', side_effect=AssertionError('unexpected FEC collection')):
                body = flow_data.flow({'root': ['committee:' + ROOT_ID], 'cycle': ['2024'],
                                       'direct': ['1']}, root=Path(directory))
            self.assertEqual(body['root']['id'], ROOT_ID)

    def test_unselected_committee_collects_once_and_draws_a_flow(self):
        with tempfile.TemporaryDirectory() as directory:
            query = {'root': ['committee:' + ROOT_ID], 'cycle': ['2024'], 'direct': ['1']}
            with patch.object(pac_flow_access.pac_evidence, 'Client'), \
                 patch.object(pac_flow_access.pac_evidence, 'collect_network', return_value=(fixture_units(), [])) as collect:
                first = flow_data.flow(query, root=Path(directory))
                second = flow_data.flow(query, root=Path(directory))
            self.assertEqual(first['root']['id'], ROOT_ID)
            self.assertEqual(second['root']['id'], ROOT_ID)
            self.assertEqual(collect.call_count, 1)
            self.assertTrue((Path(directory) / f'direct-{ROOT_ID.lower()}-2024' / 'snapshots.json').is_file())


class DonorRefunds(unittest.TestCase):
    """Lane I: SLF 2022 failed with 'Negative flow donor:WEIRMAN ROBERT -> committee:C00571703'. A donor's
    charge-backs net against its gifts in the window; one at or below zero is never drawn (never a
    negative link) and what it nets below zero is disclosed as refunded to donors."""

    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        write_fixture(cls.temp.name, units=chargeback_units())
        cls.body = produce(cls.temp.name, top=0)     # every donor named (Top 8 groups the small ones)
        cls.grouped = produce(cls.temp.name)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_no_negative_link_and_the_root_balances_gross_of_the_refunds(self):
        body, nodes = self.body, by_id(self.body)
        self.assertEqual(flow_data.contract_problems(body), [])
        self.assertEqual(flow_data.contract_problems(self.grouped), [])
        self.assertTrue(all(item['value'] > 0 for item in body['links'] + self.grouped['links']))
        for gone in ('donor:WEIRMAN ROB', 'donor:DYKE MERRITT', 'donor:REYNOLDS PETER'):
            self.assertNotIn(gone, nodes)
        self.assertEqual(link(body, 'donor:PART PAT', 'committee:' + ROOT_ID)['value'], 2000.0)
        root = nodes['committee:' + ROOT_ID]
        totals = body['totals']
        # counted in 2,047,500 (the summary nets the charge-backs) + 5,500 refunded beyond the window's
        # gifts + 100,000 prior cash = 800,000 + 550,000 + 50,000 operating + 747,500 not yet spent
        # + 5,500 refunded to donors.
        self.assertEqual((totals['counted_in'], totals['refunds_to_donors'], totals['in']), (2047500.0, 5500.0, 2153000.0))
        self.assertEqual((root['value_in'], root['value_out']), (2153000.0, 2153000.0))
        self.assertEqual(totals['operating_net'], 50000.0)
        self.assertEqual(totals['unitemized'], 1000.0)
        refunded = nodes['balance:refunded:2']
        self.assertEqual((refunded['label'], refunded['value_in']), ('Refunded to donors', 5500.0))
        self.assertEqual(refunded['details'], ['$5,500.00 in negative receipts from 2 contributors beyond their receipts '
                                               'in this window: Merritt Dyke $5,000.00 (line 11AI), Rob Weirman $500.00 (line 11AI)',
                                               '1 contributor netted to zero in this window (not drawn)'])
        self.assertIn('Refunds to donors $5,500.00: 2 contributors whose receipts in this window net below zero '
                      '(negative entries as filed); 1 netted to zero; not drawn as donors',
                      [note['text'] for note in body['notes']])

    def test_the_contract_example_is_this_producer_s_answer(self):
        example = next(item for item in contract()['examples'] if item['producer'] == 'fixture-refunds')
        self.assertEqual(json.loads(json.dumps(self.grouped)), example['body'])
        control = next(item for item in contract()['negative_controls'] if item['why'].startswith('lane I'))
        self.assertTrue(any(control['expect'] in problem for problem in flow_data.contract_problems(control['payload'])))

    def test_the_base_scenario_is_unchanged(self):
        with tempfile.TemporaryDirectory() as directory:
            write_fixture(directory)
            base = produce(directory)
        self.assertEqual((base['totals']['in'], base['totals']['refunds_to_donors']), (2151000.0, 0.0))
        self.assertNotIn('balance:refunded:2', by_id(base))

    def test_negative_control_the_old_ledger_draws_a_negative_donor_and_fails(self):
        # Teeth: put the refunded donors back as donors, as before the fix; the flow cannot be drawn.
        original = flow_data.ledger

        def old(key, unit, network):
            item = original(key, unit, network)
            if item is None:
                return item
            for entry in item['refunded_donors']:
                item['donors'][entry['key']] = {'key': entry['key'], 'names': {entry['name']: -entry['amount']},
                                                'amount': -entry['amount'], 'rows': entry['rows'], 'entity': entry['entity'],
                                                'committee': '', 'contribution': True, 'held': flow_data.ZERO,
                                                'held_kinds': set()}
            item['intake'] = item['counted_in']
            item['refunds_to_donors'] = flow_data.ZERO
            return item
        flow_data._CACHE.clear()
        flow_data.ledger = old
        try:
            with self.assertRaisesRegex(ValueError, r'Negative flow donor:(DYKE MERRITT|WEIRMAN ROB) -> committee:C90000001'):
                produce(self.temp.name)
        finally:
            flow_data.ledger = original
            flow_data._CACHE.clear()


def saved_units_in(value):
    """Every saved unit (a dict carrying the three schedules and a summary) reachable from value."""
    found, seen, stack = [], set(), [value]
    while stack:
        item = stack.pop()
        if id(item) in seen:
            continue
        seen.add(id(item))
        if isinstance(item, dict):
            if {'receipts', 'disbursements', 'expenditures', 'totals'} <= set(item):
                found.append(item)
            stack.extend(item.values())
        elif isinstance(item, (list, tuple, set)):
            stack.extend(item)
    return found


class PreparedFlow(unittest.TestCase):
    """Lane P1 (performance only): the reconciliation's evidence-network side is exactly
    downstream.snapshot_view's, without reading the snapshot file a second time; the cache keeps no
    saved unit; one request prepares a flow however many ask for it at once; and every Top N grouped
    from that one preparation is the answer a fresh preparation gives."""

    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        write_fixture(cls.temp.name)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()
        flow_data._CACHE.clear()

    def assert_the_evidence_network_s_view(self, identifier, committee, cycle, root=None):
        built = flow_data.committee_flow(committee, cycle, identifier=identifier, root=root)
        mine = flow_data.network_view(built)
        theirs = downstream.snapshot_view(identifier, committee, cycle, root)
        # snapshot_view names each route's committees; the reconciliation reads neither those nor
        # anything else beyond amounts, cycles, committees and terminal flags.
        rows = [{key: value for key, value in row.items() if key != 'route_names'} for row in theirs['relationships']]
        self.assertEqual(json.dumps(mine['relationships']), json.dumps(rows))
        self.assertEqual(json.dumps(mine['edges']), json.dumps(theirs['edges']))
        self.assertTrue(flow_data.reconcile(built, mine)['equal'])

    def test_the_fixture_s_view_is_the_evidence_network_s(self):
        self.assert_the_evidence_network_s_view('flow-fixture', ROOT_ID, '2024', Path(self.temp.name))

    @unittest.skipUnless((SAVED / FAIRSHAKE / 'snapshots.json').exists() and (SAVED / SLF / 'snapshots.json').exists(),
                         'the saved FAIRSHAKE and Senate Leadership Fund snapshots are not on this machine')
    def test_the_saved_snapshots_views_are_the_evidence_network_s(self):
        for identifier, committee, cycle in ((FAIRSHAKE, 'C00835959', '2024'), (FAIRSHAKE, 'C00835959', '2026'),
                                             (SLF, 'C00571703', '2022'), (SLF, 'C00571703', '2024'),
                                             (SLF, 'C00571703', '2026')):
            with self.subTest(snapshot=identifier, cycle=cycle):
                self.assert_the_evidence_network_s_view(identifier, committee, cycle)

    def test_one_preparation_serves_every_top_and_keeps_no_saved_unit(self):
        flow_data._CACHE.clear()
        original, calls, bodies = flow_data._prepare, [], {}

        def counted(*args, **kwargs):
            calls.append(args[:2])
            time.sleep(0.3)            # the preparation is still running when the other requests arrive
            return original(*args, **kwargs)

        def ask(top):
            bodies[top] = json.dumps(produce(self.temp.name, top=top))
        flow_data._prepare = counted
        try:
            threads = [threading.Thread(target=ask, args=(top,)) for top in (8, 3, 0)]
            for thread in threads:
                thread.start()
            for thread in threads:
                thread.join()
            ask(2)
        finally:
            flow_data._prepare = original
        self.assertEqual(calls, [(ROOT_ID, '2024')])
        self.assertEqual(len(flow_data._CACHE), 1)
        self.assertEqual(saved_units_in(list(flow_data._CACHE.values())), [])
        # Teeth: the walk the preparation came from does hold the saved units.
        self.assertTrue(saved_units_in(flow_data.committee_flow(ROOT_ID, '2024', identifier='flow-fixture',
                                                                root=Path(self.temp.name))))
        # Each Top N from the one preparation is what a fresh preparation answers.
        for top, body in bodies.items():
            flow_data._CACHE.clear()
            self.assertEqual(body, json.dumps(produce(self.temp.name, top=top)), top)
        flow_data._CACHE.clear()


class CycleWarm(unittest.TestCase):
    """Lane P1 step 2: after a cold flow, the snapshot's other cycles for the committee are prepared on one
    daemon thread when the console turns it on (server.py, on the console that warms). Off by default, off
    with MENTAT_FLOW_WARM=0, and skipped below the memory floor; what it prepares is what a click would."""

    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        units = fixture_units()
        earlier = copy.deepcopy(units[0])
        earlier['cycle'] = 2022           # the same root, a second cycle in the same snapshot file
        write_fixture(cls.temp.name, units + [earlier], identifier='flow-cycles')

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def setUp(self):
        flow_data._CACHE.clear()

    def tearDown(self):
        flow_data.WARM_CYCLES = False
        flow_data._CACHE.clear()

    def ask(self, cycle):
        return flow_data.committee_payload(ROOT_ID, cycle, identifier='flow-cycles', root=Path(self.temp.name))

    @staticmethod
    def settle():
        for thread in [thread for thread in threading.enumerate() if thread.name == 'mentat-flow-warm']:
            thread.join(60)

    def test_it_is_off_unless_the_console_turns_it_on(self):
        self.assertFalse(flow_data.WARM_CYCLES)
        self.ask('2024')
        self.settle()
        self.assertEqual(len(flow_data._CACHE), 1)

    def test_a_cold_flow_prepares_the_other_cycle_and_the_switch_reuses_it(self):
        flow_data.WARM_CYCLES = True
        original, calls = flow_data._prepare, []

        def counted(*args, **kwargs):
            calls.append(str(args[1]))
            return original(*args, **kwargs)
        flow_data._prepare = counted
        try:
            self.ask('2024')
            self.settle()
            switched = json.dumps(self.ask('2022'))
            self.ask('2024')
        finally:
            flow_data._prepare = original
        self.assertEqual(calls, ['2024', '2022'])       # the switch, and the way back, prepared nothing
        flow_data.WARM_CYCLES = False
        flow_data._CACHE.clear()
        self.assertEqual(switched, json.dumps(self.ask('2022')))

    def test_the_switch_and_the_memory_floor_turn_it_off(self):
        flow_data.WARM_CYCLES = True
        with patch.dict(os.environ, {'MENTAT_FLOW_WARM': '0'}):
            self.ask('2024')
            self.settle()
        self.assertEqual(len(flow_data._CACHE), 1)
        flow_data._CACHE.clear()
        with patch.object(flow_data.drawer_sidecars, 'free_physical_bytes', return_value=0):
            self.ask('2024')
            self.settle()
        self.assertEqual(len(flow_data._CACHE), 1)


@unittest.skipUnless((SAVED / FAIRSHAKE / 'snapshots.json').exists(), 'the saved FAIRSHAKE snapshot is not on this machine')
class Fairshake2024(unittest.TestCase):
    """Board #44's real figures, to the cent (packet section 4)."""

    @classmethod
    def setUpClass(cls):
        cls.body = flow_data.committee_payload('C00835959', '2024', identifier=FAIRSHAKE)

    def test_the_identity_holds_to_the_cent(self):
        totals, nodes = self.body['totals'], by_id(self.body)
        self.assertEqual(flow_data.contract_problems(self.body), [])
        self.assertEqual(totals['in'], 217423880.52)
        self.assertEqual(link(self.body, 'committee:C00835959', 'committee:C00836221')['value'], 54550000.00)
        self.assertEqual(link(self.body, 'committee:C00835959', 'committee:C00848440')['value'], 31700000.00)
        self.assertEqual(totals['own_candidates'], 40661820.61)
        self.assertEqual(totals['operating_net'], 3267861.97)
        self.assertEqual(totals['not_yet_spent'], {'cash': 64244197.94, 'held': 23000000.0, 'total': 87244197.94,
                                                   'held_kinds': ['bitcoin']})
        self.assertEqual(round(54550000.00 + 31700000.00 + 40661820.61 + totals['operating_net'] + 87244197.94, 2),
                         217423880.52)
        self.assertEqual(totals['in_kind'], {'booked': 64660404.58, 'proceeds': 41660409.58, 'held': 23000000.0,
                                             'excess_proceeds': 5.0})
        self.assertEqual(nodes['bucket:own-ads']['value_in'], 40661820.61)

    def test_the_board_s_columns(self):
        nodes = by_id(self.body)
        donors = {nodes[item['source']]['label']: item['value'] for item in self.body['links']
                  if item['target'] == 'committee:C00835959' and nodes[item['source']]['stage'] == 0}
        self.assertEqual(donors, {'Coinbase': 68499995.00, 'AH Capital Management': 67000000.01, 'Ripple Labs': 45000000.00,
                                  'Jump Crypto': 25000000.00, 'Cameron Winklevoss': 2450000.00,
                                  'Tyler Winklevoss': 2450000.00, 'All other contributors': 7023885.51})
        self.assertEqual(nodes['donor:COINBASE']['held'], 23000000.0)
        jobs, progress = nodes['committee:C00836221'], nodes['committee:C00848440']
        self.assertEqual(jobs['share_from_upstream'][0]['share'], 0.904269)
        self.assertEqual(progress['share_from_upstream'][0]['share'], 0.833421)
        self.assertEqual(nodes['stub:external:committee:C00836221:C00836221']['value_out'], 5775000.01)
        self.assertEqual(nodes['stub:external:committee:C00848440:C00848440']['value_out'], 6275255.01)
        # Each candidate, and the Other node, holds exactly what the evidence network measures for
        # the drawn committees (the measurement's own figures, whatever F12 does to them).
        view = downstream.snapshot_view(FAIRSHAKE, 'C00835959', '2024')
        drawn = set(self.body['reconciliation']['committees'])
        measured = {}
        for row in view['relationships']:
            if row['committee'] in drawn and row['amount'] > 0:
                word = {'Support': 'support', 'Oppose': 'oppose'}.get(row['direction'], 'contribution'
                                                                       if row['direction'] == 'Contribution / Other' else 'unknown')
                key = f"candidate:{row['candidate_id'] or 'unidentified'}:{word}"
                measured[key] = round(measured.get(key, 0) + row['amount'], 2)
        landed = [node for node in self.body['nodes'] if node['stage'] == 3 and node['kind'] == 'candidate']
        for node in landed:
            self.assertEqual(node['value_in'], measured[node['id']], node['label'])
        other = nodes['other:3']
        self.assertEqual(other['label'], 'Other candidates')
        self.assertEqual(other['value_in'], round(sum(measured.values()) - sum(node['value_in'] for node in landed), 2))
        self.assertNotEqual(round(other['value_in'] / 1e6, 1), 55.0)      # the board's $55.0M was wrong
        self.assertEqual(self.body['reconciliation']['candidates']['flow'],
                         self.body['reconciliation']['candidates']['evidence_network'])
        self.assertEqual([node['label'] for node in self.body['nodes'] if node['stage'] == 3 and node['kind'] == 'balance'],
                         ['Kept by DAJ and PP'])
        self.assertEqual([note['text'] for note in self.body['notes']][:2], [
            'Donated crypto counted once, at its value when given (FAIRSHAKE booked $64.7M in and out again)',
            'routes through Defend American Jobs and Protect Progress show where pooled money went'])

    def test_the_recorded_example_is_this_producer_s_answer(self):
        # The root's side of the recorded example is the producer's to the cent; the downstream side
        # is the measurement's (asserted above against it), so a dated walk may move it.
        recorded = next(example for example in contract()['examples'] if example.get('producer') == 'fairshake-2024')
        self.assertEqual(flow_data.contract_problems(recorded['body']), [])
        mine = json.loads(json.dumps(self.body))
        for field in ('version', 'root', 'cycle', 'window', 'basis', 'stages'):
            self.assertEqual(mine[field], recorded['body'][field], field)
        root_totals = ('in', 'counted_in', 'donors', 'unitemized', 'from_network', 'cash_before', 'transfers_out',
                       'own_candidates', 'operating_net', 'not_yet_spent', 'in_kind', 'refunds_netted')
        self.assertEqual({key: mine['totals'][key] for key in root_totals},
                         {key: recorded['body']['totals'][key] for key in root_totals})
        near_root = lambda body: sorted((node['id'], node['value_in'], node['value_out']) for node in body['nodes']
                                        if node['stage'] <= 1 and not node.get('stub'))
        self.assertEqual(near_root(mine), near_root(recorded['body']))


@unittest.skipUnless((SAVED / SLF / 'snapshots.json').exists(), 'the saved Senate Leadership Fund snapshot is not on this machine')
class SenateLeadershipFund2024(unittest.TestCase):
    """The loops lane D measured: American Crossroads and America Fund send money back."""

    def test_loops_are_stubs_and_every_node_balances(self):
        body = flow_data.committee_payload('C00571703', '2024', identifier=SLF)
        self.assertEqual(flow_data.contract_problems(body), [])
        nodes = by_id(body)
        root = 'committee:C00571703'
        self.assertEqual(nodes[f'stub:return_in:{root}:C00487363']['value_out'], 1305577.91)
        self.assertEqual(nodes['stub:return:committee:C00487363:C00571703']['value_in'], 1300000.00)
        walk = downstream.snapshot_view(SLF, 'C00571703', '2024')
        # F12: Conservative Americans PAC paid America Fund before SLF's money reached it, so the
        # branch is out of the network; America Fund's $5M to SLF is an ordinary donor now.
        self.assertFalse(any(edge['source'] == 'C00820308' and edge['destination'] == 'C00782664'
                             for edge in walk['edges']))
        self.assertNotIn(f'stub:return_in:{root}:C00782664', nodes)
        whole = flow_data.committee_payload('C00571703', '2024', identifier=SLF, top=0)
        self.assertEqual(link(whole, 'donor:AMERICA FUND', root)['value'], 5000000.00)
        # The money each drawn committee paid before SLF's arrived is disclosed beside it, to the
        # measurement's own gate, and none of it is a flow.
        gated = [item for item in body['links'] if item['kind'] == 'uncounted']
        self.assertTrue(gated)
        self.assertTrue(all(nodes[item['target']]['stub'] == 'before_gate' for item in gated))
        self.assertEqual(body['totals']['before_gate'], round(sum(item['value'] for item in gated), 2))
        self.assertTrue(body['reconciliation']['equal'])


def carry_problems(body, rule):
    """Every drawn committee's gate source is the measurement's, and every carried-over one says
    so in the measurement's words, on the node and on its starting cash."""
    problems = []
    sources = {gate['committee']: gate.get('gate_source') for gate in rule['gates'] if str(gate['cycle']) == body['cycle']}
    opened = {entry['committee']: entry for entry in (rule.get('carryover') or {}).get('opened') or []
              if str(entry['cycle']) == body['cycle']}
    nodes = by_id(body)
    for node in body['nodes']:
        if node['kind'] != 'committee' or node.get('stub') or node['committee_id'] not in sources:
            continue
        ident = node['committee_id']
        if node.get('gate_source') != sources[ident]:
            problems.append(f'{ident}: gate source {node.get("gate_source")!r}, the measurement says {sources[ident]!r}')
        if sources[ident] == 'carryover':
            label = opened[ident]['label']
            if (node.get('carryover') or {}).get('label') != label:
                problems.append(f'{ident}: no carryover label')
            cash = nodes.get(f'stub:prior_cash:{node["id"]}:{ident}')
            if cash is not None and cash.get('carryover') != label:
                problems.append(f'{ident}: its starting cash does not say it was carried over')
    return problems


@unittest.skipUnless((SAVED / SLF / 'snapshots.json').exists() and (SAVED / FAIRSHAKE / 'snapshots.json').exists(),
                     'the saved FAIRSHAKE and Senate Leadership Fund snapshots are not on this machine')
class Carryover(unittest.TestCase):
    """Carryover (option (b), landed by lane B 2026-09-26): the flow shows where each gate came
    from in the measurement's own words (date_rule), and never derives a gate itself."""

    def test_senate_leadership_fund_2024_names_its_carried_over_committees(self):
        body = flow_data.committee_payload('C00571703', '2024', identifier=SLF)
        rule = downstream.snapshot_view(SLF, 'C00571703', '2024')['date_rule']
        self.assertEqual(flow_data.contract_problems(body), [])
        self.assertEqual(carry_problems(body, rule), [])
        crossroads = by_id(body)['committee:C00487363']
        self.assertEqual((crossroads['gate_source'], crossroads['carryover']['label']),
                         ('carryover', 'gate opened by 2021-2022 transfers'))
        self.assertTrue(crossroads['carryover']['opened_by'])
        carried = rule['carryover']
        self.assertEqual(body['carryover']['status'], carried['status'])
        opened = [entry for entry in carried['opened'] if str(entry['cycle']) == '2024']
        texts = [note['text'] for note in body['notes']]
        self.assertIn(f"{len(opened)} committees counted from the cycle's first day: {opened[0]['label']}", texts)
        if carried['status'] == 'partial':
            self.assertIn(flow_data.CARRYOVER_PARTIAL, texts)
        self.assertTrue(body['reconciliation']['equal'])

    def test_fairshake_2026_carries_protect_progress_s_spending_back_in(self):
        body = flow_data.committee_payload('C00835959', '2026', identifier=FAIRSHAKE, top=0)
        rule = downstream.snapshot_view(FAIRSHAKE, 'C00835959', '2026')['date_rule']
        self.assertEqual(flow_data.contract_problems(body), [])
        self.assertEqual(carry_problems(body, rule), [])
        self.assertTrue(body['reconciliation']['equal'])
        progress = by_id(body)['committee:C00848440']
        self.assertEqual(progress['carryover']['label'], 'gate opened by 2023-2024 transfers')
        # Lane B's measured returns (out/compare12c.txt): Protect Progress's 2025 spending.
        people = {node['label']: node['value_in'] for node in body['nodes'] if node['kind'] == 'candidate'}
        self.assertEqual(people['James Walkinshaw'], 1012792.00)
        self.assertEqual(people['Stella Pekarsky'], 10241.45)

    def test_fairshake_2024_says_its_carryover_is_unknown_and_keeps_its_identity(self):
        body = flow_data.committee_payload('C00835959', '2024', identifier=FAIRSHAKE)
        self.assertEqual(body['carryover']['status'], 'unknown')
        self.assertIn(flow_data.CARRYOVER_UNKNOWN, [note['text'] for note in body['notes']])
        self.assertEqual(body['totals']['in'], 217423880.52)
        self.assertEqual(carry_problems(body, downstream.snapshot_view(FAIRSHAKE, 'C00835959', '2024')['date_rule']), [])

    def test_senate_leadership_fund_2026_at_top_8_groups_a_carried_over_committee_cleanly(self):
        # The sweep's one failure: a carried-over committee grouped into Other left its "carried over"
        # mark on the cash stub merged into the Other node. Teeth: put the mark back and the rule fires.
        body = flow_data.committee_payload('C00571703', '2026', identifier=SLF)
        self.assertEqual(flow_data.contract_problems(body), [])
        merged = [node for node in body['nodes'] if node.get('stub') and str(node.get('attach', '')).startswith('other:')]
        self.assertTrue(merged, 'the Top 8 answer groups committees with stubs into Other')
        self.assertTrue(all('carryover' not in node and 'gate' not in node for node in merged))
        members = [member for node in body['nodes'] if node.get('bucket') == 'other' for member in node.get('members') or []]
        self.assertTrue(any(member.get('carryover') == 'gate opened by 2023-2024 transfers' for member in members))
        broken = copy.deepcopy(body)
        cash = next(node for node in broken['nodes'] if node.get('stub') == 'prior_cash'
                    and str(node.get('attach', '')).startswith('other:'))
        cash['carryover'] = 'gate opened by 2023-2024 transfers'
        self.assertTrue(any('says carried over but its committee is not' in problem
                            for problem in flow_data.contract_problems(broken)))

    def test_negative_control_a_lost_label_or_a_gate_of_its_own_is_caught(self):
        body = flow_data.committee_payload('C00571703', '2024', identifier=SLF)
        rule = downstream.snapshot_view(SLF, 'C00571703', '2024')['date_rule']
        unlabelled = copy.deepcopy(body)
        del by_id(unlabelled)['committee:C00487363']['carryover']
        self.assertTrue(carry_problems(unlabelled, rule))
        rederived = copy.deepcopy(body)
        by_id(rederived)['committee:C00487363']['gate_source'] = 'in-cycle'
        self.assertTrue(carry_problems(rederived, rule))


@unittest.skipUnless(CRYPTO.exists() and (SAVED / FAIRSHAKE / 'snapshots.json').exists(),
                     'the Crypto workbook or the saved FAIRSHAKE snapshot is not on this machine')
class CryptoIndustry2024(unittest.TestCase):
    """Board #45 on Reported: companies, first recipients, passed to, outcome."""

    @classmethod
    def setUpClass(cls):
        cls.body = flow_data.industry_payload('Crypto', '2024', [CRYPTO], top=8)

    def test_company_money_counts_once_and_the_flow_reconciles(self):
        body, nodes = self.body, by_id(self.body)
        self.assertEqual(flow_data.contract_problems(body), [])
        self.assertTrue(body['reconciliation']['equal'])
        self.assertEqual([stage['label'] for stage in body['stages']], ['Companies', 'First recipient', 'Passed to', 'Outcome'])
        fairshake = nodes['committee:C00835959']
        self.assertEqual((fairshake['stage'], fairshake['value_in']), (1, 217423880.52))
        coinbase = link(body, 'company:Coinbase', 'committee:C00835959')
        outside = nodes['stub:external:committee:C00835959:C00835959']['value_out']
        self.assertEqual(round(coinbase['value'] + outside, 2), 217423880.52)
        self.assertLess(coinbase['value'], 125511179.52)   # the workbook's figure counts sale proceeds again
        self.assertEqual(nodes['committee:C00836221']['stage'], 2)   # passed to, though Coinbase also gave directly
        self.assertEqual(body['totals']['counted_from_filings'] < body['totals']['workbook_for_those_committees'], True)
        self.assertTrue(any(note['key'] == 'counted_once' for note in body['notes']))

    def test_the_recorded_example_is_this_producer_s_answer(self):
        recorded = next(example for example in contract()['examples']
                        if example.get('producer') == 'crypto-2024-industry' and example['status'] == 200)
        self.assertEqual(flow_data.contract_problems(recorded['body']), [])
        mine = json.loads(json.dumps(self.body))
        companies = lambda body: sorted((item['source'], item['target'], item['value']) for item in body['links']
                                        if item['source'].startswith('company:'))
        self.assertEqual(companies(mine), companies(recorded['body']))
        self.assertEqual(mine['stages'], recorded['body']['stages'])

    def test_attributed_is_refused_without_an_estimate_bundle(self):
        with self.assertRaises(flow_data.FlowRefused) as caught:
            flow_data.flow({'root': ['industry'], 'industry': ['Crypto'], 'cycle': ['2024'], 'basis': ['attributed']},
                           workbooks=lambda label: [CRYPTO])
        self.assertEqual(caught.exception.reason, 'no_estimate_bundle')


BANKING = Path(r'C:\Users\pgoff\OneDrive\Desktop\Mentat Complete Files\Banking Industry Spending Overview.xlsx')


@unittest.skipUnless(BANKING.exists(), 'the Banking workbook is not on this machine')
class IndustryRefunds(unittest.TestCase):
    """The industry sweep: Banking 2021-2022 answered 409 'bucket:candidates does not balance' ($265.00).
    A company-candidate pair that nets below zero is not drawn, and the bucket passes on what it drew."""

    @classmethod
    def setUpClass(cls):
        # The Banking workbook's Transfer Paths sheet is a partial display, so this flow reads the full
        # report source pinned to the workbook (graph_data.transfer_source): a payload and its manifest
        # under the gitignored mentat/.cache/industry-report. A checkout without them cannot build the
        # flow, so the case skips and names what is missing instead of erroring (lane Z's clean export).
        import graph_data
        missing = [path for path in graph_data.transfer_source(BANKING) if path is not None and not path.is_file()]
        if missing:
            raise unittest.SkipTest("the Banking workbook's pinned full report source is not on this machine: "
                                    + ', '.join(str(path) for path in missing))
        cls.built = flow_data.industry_flow('Banking', '2022', [BANKING], top=8)

    def test_banking_2022_is_drawn_and_every_node_balances(self):
        body = flow_data.industry_payload('Banking', '2022', [BANKING], top=8)
        self.assertEqual(flow_data.contract_problems(body), [])
        self.assertTrue(all(item['value'] > 0 for item in body['links']))
        refunds = sum(pair['amount'] for pair in self.built['refunds_back'][True])
        self.assertGreaterEqual(refunds, Decimal('265.00'))   # the sweep's imbalance; undrawn pairs may add to it
        self.assertTrue(any(note['text'].startswith(f'Refunds from candidates {flow_data.exact_money(refunds)}:')
                            for note in body['notes']))
        nodes = by_id(body)
        bucket = nodes['bucket:candidates']
        self.assertEqual(bucket['value_in'], bucket['value_out'])

    def test_negative_control_passing_on_the_net_does_not_balance(self):
        # Teeth: the old rule sent each candidate's NET total on; the bucket then takes in $265.00 more
        # than it passes on, and the flow is refused for conservation.
        graph = copy.deepcopy(self.built['graph'])
        refunds = sum(pair['amount'] for pair in self.built['refunds_back'][True])
        out = next(link for link in graph.links.values() if link['source'] == 'bucket:candidates')
        out['value'] -= refunds
        nodes = OrderedDict((key, dict(node)) for key, node in graph.nodes.items())
        links = [dict(link) for link in graph.links.values()]
        flow_data.measure(nodes, links)
        self.assertTrue(any(ident == 'bucket:candidates' for ident, _in, _out in flow_data.conservation_errors(nodes)))


def start_console():
    import server as app

    class Console(app.Console):
        def __init__(self, directory):
            self.root = Path(directory)
            self.report_directory = Path(directory)
            self.token = 'flow-test-token'
            self.report_cache, self.processes, self.lock = {}, {}, threading.Lock()

    temp = tempfile.TemporaryDirectory()
    httpd = app.ThreadingHTTPServer(('127.0.0.1', 0), app.Handler)
    httpd.app = Console(temp.name)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    return httpd, temp


def get(port, query):
    try:
        with urllib.request.urlopen(f'http://127.0.0.1:{port}/api/flow?{query}', timeout=60) as answer:
            return answer.status, json.loads(answer.read())
    except urllib.error.HTTPError as error:
        return error.code, json.loads(error.read())


@unittest.skipUnless((SAVED / FAIRSHAKE / 'snapshots.json').exists(), 'the saved FAIRSHAKE snapshot is not on this machine')
class Route(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.httpd, cls.temp = start_console()
        cls.port = cls.httpd.server_address[1]

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown()
        cls.httpd.server_close()
        cls.temp.cleanup()

    def test_the_route_answers_reports_and_refuses(self):
        base = f'root=committee:C00835959&cycle=2024&id={FAIRSHAKE}'
        status, body = get(self.port, base + '&top=8&depth=3')
        self.assertEqual(status, 200)
        self.assertEqual(flow_data.contract_problems(body), [])
        self.assertEqual(body['totals']['in'], 217423880.52)
        status, body = get(self.port, base + '&basis=attributed')
        self.assertEqual((status, body['reason']), (409, 'no_estimate_bundle'))
        status, body = get(self.port, 'root=committee:C00835959&cycle=2023')
        self.assertEqual(status, 400)

    def test_the_route_refuses_a_flow_that_disagrees(self):
        # The route reads the evidence network's side through flow_data.network_view (lane P1), which
        # PreparedFlow holds to downstream.snapshot_view's relationships and edges.
        original = flow_data.network_view

        def disagreeing(*args, **kwargs):
            view = original(*args, **kwargs)
            view['relationships'][0]['amount'] += 1.0
            return view
        flow_data._CACHE.clear()
        flow_data.network_view = disagreeing
        try:
            status, body = get(self.port, f'root=committee:C00835959&cycle=2024&id={FAIRSHAKE}')
        finally:
            flow_data.network_view = original
            flow_data._CACHE.clear()
        self.assertEqual((status, body['reason'], body['reconciliation']['equal']), (409, 'reconciliation', False))


class CommitteeNames(unittest.TestCase):
    """B4 (lane Z, 2026-10-15): a committee reached only as a recipient is named from the FEC committee master
    on disk, never by its bare ID and never by a name made up for it. Nothing is fetched."""

    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.master = Path(cls.temp.name) / 'master'
        (cls.master / 'cm26').mkdir(parents=True)
        (cls.master / 'cm26' / 'cm.txt').write_text(
            f'{ONWARD_ID}|ONWARD POLITICAL ACTION COMMITTEE|TREASURER, A|1 MAIN ST||CITY|DC|20001|B|Q||M|C||\n'
            'C90000099|SOMEONE ELSE PAC|T|1 MAIN ST||CITY|DC|20001|B|Q||M|C||\n', encoding='ascii')
        cls.saved = flow_data.MASTER_ROOT
        flow_data.MASTER_ROOT = cls.master

    @classmethod
    def tearDownClass(cls):
        flow_data.MASTER_ROOT = cls.saved
        cls.temp.cleanup()

    def test_names_come_from_the_saved_data_first_then_the_master_and_otherwise_none_is_made_up(self):
        names = flow_data.Names({'C90000010': 'FILED NAME PAC', 'C90000011': 'C90000011'})
        self.assertEqual(flow_data.committee_label(names, 'C90000010'), 'Filed Name PAC')
        self.assertEqual(flow_data.committee_label(names, ONWARD_ID), 'Onward Political Action Committee')     # master
        self.assertEqual(names.get(ONWARD_ID), 'ONWARD POLITICAL ACTION COMMITTEE')
        # An ID the ledger fell back to is no name; with none on file anywhere it is "Unnamed committee".
        for ident in ('C90000011', 'C90000012'):
            self.assertFalse(names.named(ident))
            self.assertEqual(flow_data.committee_label(names, ident), 'Unnamed committee')
            self.assertEqual(flow_data.committee_tag(names, ident), f'Unnamed committee {ident}')
            self.assertEqual(flow_data.unnamed_fields(names, ident), {'unnamed': True})
        self.assertEqual(flow_data.unnamed_fields(names, ONWARD_ID), {})

    def test_the_flow_says_where_money_went_by_name_and_by_ID_when_there_is_no_name(self):
        def stub_text(directory, with_name):
            folder = Path(directory) / 'flow-fixture'
            write_fixture(directory)
            evidence = json.loads((folder / 'evidence.json').read_text(encoding='ascii'))
            evidence['transfers'] = [{'destination': ONWARD_ID, 'destination_name': 'ONWARD PAC'}] if with_name else []
            (folder / 'evidence.json').write_text(json.dumps(evidence), encoding='ascii')
            flow_data._CACHE.clear()
            body = produce(directory)
            return by_id(body)[f'stub:onward:committee:{BRIDGE_ID}:{ONWARD_ID}']['label']
        with tempfile.TemporaryDirectory() as directory:
            self.assertEqual(stub_text(directory, True), 'passed to Onward PAC', 'the saved data\'s own name wins')
            self.assertEqual(stub_text(directory, False), 'passed to Onward Political Action Committee', 'the master names it')
            flow_data.MASTER_ROOT = Path(directory) / 'no-master'
            try:
                flow_data._CACHE.clear()
                text = stub_text(directory, False)
            finally:
                flow_data.MASTER_ROOT = self.master
            self.assertEqual(text, f'passed to Unnamed committee {ONWARD_ID}')
            self.assertNotIn('ONWARD', text.upper().replace(ONWARD_ID, ''), 'no name was made up')
        flow_data._CACHE.clear()

    def test_a_committee_is_named_from_the_master_file_of_the_flows_cycle_then_the_newest(self):
        """657 committee IDs differ between cm22 and cm26 (C00728980 is "TIPPINGPOINT 2020", then "DEMOCRACY SOLVES
        IT"): the flow of a cycle names a committee as that cycle's file does, and falls back to the newest file."""
        with tempfile.TemporaryDirectory() as directory:
            master = Path(directory)
            line = '{}|{}|T|1 MAIN ST||CITY|DC|20001|B|Q||M|C||\n'
            for folder, text in (('cm22', line.format('C00728980', 'TIPPINGPOINT 2020')),
                                 ('cm24', line.format('C00728980', 'IN BETWEEN')),
                                 ('cm26', line.format('C00728980', 'DEMOCRACY SOLVES IT')
                                  + line.format('C90000077', 'ONLY IN THE NEWEST'))):
                (master / folder).mkdir()
                (master / folder / 'cm.txt').write_text(text, encoding='ascii')
            flow_data.MASTER_ROOT = master
            try:
                for cycle, name, year in ((2022, 'TIPPINGPOINT 2020', '2022'), ('2022', 'TIPPINGPOINT 2020', '2022'),
                                          (2024, 'IN BETWEEN', '2024'), (2026, 'DEMOCRACY SOLVES IT', '2026'),
                                          (2028, 'DEMOCRACY SOLVES IT', '2026'), (None, 'DEMOCRACY SOLVES IT', '2026')):
                    names = flow_data.Names({}, cycle=cycle)
                    self.assertEqual((names.get('C00728980'), names.master_year('C00728980')), (name, year), cycle)
                    self.assertEqual(flow_data.master_fields(names, 'C00728980'), {'name_master': year})
                # A committee the cycle's file does not name still gets the newest file's name, and says which year.
                names = flow_data.Names({}, cycle=2022)
                self.assertEqual((names.get('C90000077'), names.master_year('C90000077')), ('ONLY IN THE NEWEST', '2026'))
                # The saved filings' own name wins over any file, and then no year is claimed.
                filed = flow_data.Names({'C00728980': 'FILED NAME'}, cycle=2022)
                self.assertEqual((filed.get('C00728980'), filed.master_year('C00728980')), ('FILED NAME', ''))
                self.assertEqual(flow_data.master_fields(filed, 'C00728980'), {})
                # No name anywhere: nothing is made up and no year is claimed.
                self.assertEqual(flow_data.master_fields(names, 'C90000555'), {})
            finally:
                flow_data.MASTER_ROOT = self.master

    def test_a_flow_asks_for_the_names_of_its_own_cycle_and_marks_a_master_name_on_the_node(self):
        with tempfile.TemporaryDirectory() as directory:
            master = Path(directory) / 'master'
            for folder, name in (('cm24', 'BRIDGE PAC IN 2024'), ('cm26', 'BRIDGE PAC IN 2026')):
                (master / folder).mkdir(parents=True)
                (master / folder / 'cm.txt').write_text(f'{BRIDGE_ID}|{name}|T|1 MAIN ST||CITY|DC|20001|B|Q||M|C||\n', encoding='ascii')
            units = fixture_units()
            units[1]['totals']['committee_name'] = BRIDGE_ID           # the saved filings give the bridge no name
            write_fixture(directory, units)
            asked = []
            real = flow_data.committee_names
            flow_data.MASTER_ROOT = master
            flow_data._CACHE.clear()
            try:
                with patch.object(flow_data, 'committee_names', side_effect=lambda cycle=None: (asked.append(cycle), real(cycle))[1]):
                    node = by_id(produce(directory))['committee:' + BRIDGE_ID]
            finally:
                flow_data.MASTER_ROOT = self.master
                flow_data._CACHE.clear()
            self.assertIn('2024', [str(cycle) for cycle in asked])
            self.assertEqual((node['name_as_filed'], node['name_master']), ('BRIDGE PAC IN 2024', '2024'))
            self.assertNotIn('unnamed', node)


class CandidateNames(unittest.TestCase):
    """A candidate node is placed by the FEC candidate master and named by the filed spellings the master agrees
    with. Coinbase PAC filed "ROSEN, JACKY" (NV) under Deb Fischer's S2NE00094; the alphabet picked ROSEN. The
    master is the legal name ("James Johnson" for Mike Johnson), so it chooses among the filed spellings and is the
    label only when none agrees. Labels and state only: the node's dollars and the reconciliation do not move."""
    CAND = 'H4XX01001'
    NODE = 'candidate:H4XX01001:contribution'
    ROWS = (('k1', 'SMITH, ANN', 'XX', 2500.0), ('k2', 'ZEDDY, ZOE', 'YY', 1000.0))

    def units(self, rows=ROWS):
        units = fixture_units()
        for unit in units:                                  # the base scenario's own SMITH rows move to another ID
            for schedule in ('disbursements', 'expenditures'):
                for row in unit[schedule]:
                    if row.get('candidate_id') == self.CAND:
                        row['candidate_id'] = 'H4XX09009'
        for sub_id, name, state, value, *party in rows:
            row = payment(sub_id, 'SMITH FOR CONGRESS', value, '2023-10-02', line='23', kind='24K',
                          committee='C90000009', candidate=self.CAND)
            row.update(candidate_name=name, candidate_office_state=state)
            if party:
                row['candidate_party'] = party[0]
            units[1]['disbursements'].append(row)
        return units

    def flow(self, master_lines=None, rows=ROWS):
        with tempfile.TemporaryDirectory() as directory:
            master = Path(directory) / 'master'
            master.mkdir()
            if master_lines is not None:
                (master / 'cn24').mkdir()
                (master / 'cn24' / 'cn.txt').write_text(master_lines, encoding='ascii')
            write_fixture(directory, self.units(rows))
            saved, flow_data.CANDIDATE_MASTER_ROOT = flow_data.CANDIDATE_MASTER_ROOT, master
            flow_data._CACHE.clear()
            try:
                return produce(directory)
            finally:
                flow_data.CANDIDATE_MASTER_ROOT = saved
                flow_data._CACHE.clear()

    def received(self, body):
        return sum(item['value'] for item in body['links'] if item['target'] == self.NODE)

    def master(self, name, state='XX'):
        return f'{self.CAND}|{name}|REP|2024|{state}|H|01|I|C|C90000009|1 MAIN ST||CITY|{state}|20001\n'

    # ------------------------------------------------------------ one rule in every view (Flow, Graph direct, Graph activity)
    def graph(self, master_lines, names, *, mode='activity', other=(), between=None, seen=None, folders=None):
        """Graph (graph_data.graph_view) over the shared fixture workbook with its candidate re-keyed to CAND and
        its filed names replaced by `names` (the PAC map's sponsor gets the first). `other` adds Downstream
        Activity rows. The candidate master is `master_lines` (cn24)."""
        import openpyxl
        import graph_data
        from graph_fixture import add_graph_sheets
        with tempfile.TemporaryDirectory() as directory:
            master = Path(directory) / 'master'
            (master / 'cn24').mkdir(parents=True)
            (master / 'cn24' / 'cn.txt').write_text(master_lines, encoding='ascii')
            for folder, text in (folders or {}).items():   # more master files: {'cn26': lines}
                (master / folder).mkdir()
                (master / folder / 'cn.txt').write_text(text, encoding='ascii')
            path = Path(directory) / 'Giving.xlsx'
            book = openpyxl.Workbook()
            add_graph_sheets(book)
            for row in book['PAC & Candidate Map'].iter_rows(min_row=6):
                if row[5].value == 'H0CA00001':
                    row[4].value, row[5].value = names[0], self.CAND
            sheet = book['Downstream Activity']
            filed = iter(names * 4)
            for row in sheet.iter_rows(min_row=6):
                if row[2].value == 'H0CA00001':
                    row[2].value, row[3].value = self.CAND, next(filed)
            for row in other:
                sheet.append(row)
            sheet['A3'] = f'{100 + len(other)} source rows condensed to {4 + len(other)} committee-candidate-activity groups'
            book.save(path)
            book.close()
            saved, flow_data.CANDIDATE_MASTER_ROOT = flow_data.CANDIDATE_MASTER_ROOT, master
            try:
                if between is not None:                 # read once, change the master, read again (the cache)
                    seen.append(graph_data.graph_view([path], {'mode': mode}))
                    os.utime(master / 'cn24' / 'cn.txt', ns=(1, 1))
                    between(master)
                return graph_data.graph_view([path], {'mode': mode})
            finally:
                flow_data.CANDIDATE_MASTER_ROOT = saved

    def graph_labels(self, body):
        """Every label Graph gives each candidate ID: the ranking, the direction sides and the details table."""
        seen = defaultdict(set)
        for row in body['candidates'] + [row for rows in body['candidates_by_direction'].values() for row in rows]:
            seen[row['id']].add(row['label'])
        for row in body['candidate_details'] + body['pac_candidate_relationships']:
            seen[row['candidate_id']].add(row['candidate'])
        return dict(seen)

    def test_one_candidate_reads_the_same_in_flow_and_in_both_graph_views(self):
        master = self.master('SMITH, ANNABELLE ROSE')
        flow_label = by_id(self.flow(master))[self.NODE]['label']
        direct = self.graph(master, ['SMITH, ANN'], mode='direct')
        activity = self.graph(master, ['SMITH, ANN', 'ZEDDY, ZOE'])
        self.assertEqual(flow_label, 'Ann Smith')
        self.assertEqual(self.graph_labels(direct)[self.CAND], {flow_label})
        self.assertEqual(self.graph_labels(activity)[self.CAND], {flow_label})
        self.assertTrue(any(row['candidate_id'] == self.CAND for row in activity['pac_candidate_relationships']))

    def test_graph_names_a_candidate_by_flows_rule_not_by_the_last_row_read(self):
        """Regression (Graph activity): the label was the last workbook row's spelling, so a mis-keyed row read
        last renamed the candidate. Now the spellings vote and the master chooses, as in Flow; amounts stay."""
        master = self.master('SMITH, ANN')
        before = self.graph(master, ['SMITH, ANN'])
        body = self.graph(master, ['SMITH, ANN', 'SMITH, ANN', 'SMITH, ANN', 'ZEDDY, ZOE'])
        self.assertEqual(self.graph_labels(body)[self.CAND], {'Ann Smith'})
        self.assertEqual([(row['id'], row['amount']) for row in body['candidates']],
                         [(row['id'], row['amount']) for row in before['candidates']])

    def test_graph_direct_view_takes_the_masters_name_when_no_filed_spelling_agrees(self):
        """Regression (Graph direct): the PAC map's sponsor text was the label as it stood."""
        body = self.graph(self.master('SMITH, ANN'), ['Candidate One'], mode='direct')
        self.assertEqual(self.graph_labels(body)[self.CAND], {'Ann Smith'})
        self.assertEqual([row['amount'] for row in body['candidates']], [90])

    def test_a_shared_first_name_with_another_surname_stays_another_person_in_every_view(self):
        """JOHNSON, ANN filed under SMITH, ANN's ID neither names the node nor merges with the real Ann Johnson."""
        johnson = 'H4XX02002|JOHNSON, ANN|DEM|2024|XX|H|02|C|C|C90000010|1 MAIN ST||CITY|XX|20001\n'
        master = self.master('SMITH, ANN') + johnson
        node = by_id(self.flow(master, rows=(('k1', 'JOHNSON, ANN', 'XX', 2500.0),)))[self.NODE]
        self.assertEqual(node['label'], 'Ann Smith')
        body = self.graph(master, ['JOHNSON, ANN'],
                          other=(['2025-2026', 'C00000003', 'H4XX02002', 'JOHNSON, ANN', 'D', 'Support', 50],))
        labels = self.graph_labels(body)
        self.assertEqual((labels[self.CAND], labels['H4XX02002']), ({'Ann Smith'}, {'Ann Johnson'}))
        self.assertEqual(flow_data.CandidateNames({}, masters=[('2024', {})]).resolve('H4XX02002', 'JOHNSON, ANN')['label'],
                         'Ann Johnson')

    def test_an_unfiltered_one_cycle_graph_reads_its_own_cycles_master_before_a_newer_one(self):
        """The fixture report is 2025-2026: with no cycle chosen, cn26 names the candidate before the newer cn28,
        as Flow reads its cycle's file first."""
        body = self.graph(self.master('JONES, PAT'), ['Candidate One'], mode='direct',
                          folders={'cn26': self.master('SMITH, ANN'), 'cn28': self.master('ZEDDY, ZOE')})
        self.assertEqual(self.graph_labels(body)[self.CAND], {'Ann Smith'})

    def test_the_graph_cache_reads_a_changed_candidate_master_again(self):
        changed = []
        body = self.graph(self.master('JONES, PAT'), ['Candidate One'], mode='direct',
                          between=lambda folder: (folder / 'cn24' / 'cn.txt').write_text(self.master('SMITH, ANN'), encoding='ascii'),
                          seen=changed)
        self.assertEqual(self.graph_labels(changed[0])[self.CAND], {'Pat Jones'})
        self.assertEqual(self.graph_labels(body)[self.CAND], {'Ann Smith'})

    # ------------------------------------------------------------ party from the cycle's FEC candidate master (D-343)
    ROSEN_FIRST = (('k2', 'ZEDDY, ZOE', 'YY', 1000.0, 'DEM'), ('k1', 'SMITH, ANN', 'XX', 2500.0, 'REP'))

    def test_the_rosen_fischer_case_takes_the_masters_party_not_the_first_filing(self):
        """The first filing under the ID is the mis-keyed one (ROSEN, DEM under Fischer's REP ID): the node is the
        master's party, unflagged, and its dollars do not move."""
        body = self.flow(self.master('SMITH, ANN'), rows=self.ROSEN_FIRST)
        node = by_id(body)[self.NODE]
        self.assertEqual(node['party'], 'R')
        self.assertNotIn('party_source', node)
        self.assertEqual(self.received(body), 3500.0)
        self.assertEqual(by_id(self.flow(None, rows=self.ROSEN_FIRST))[self.NODE]['party'], 'D',
                         'without a master the first filing still answers (and is flagged)')

    def test_a_candidate_missing_from_the_cycles_master_keeps_the_filed_party_and_is_flagged(self):
        other = self.master('SMITH, ANN').replace(self.CAND, 'H4XX09999')
        node = by_id(self.flow(other, rows=self.ROSEN_FIRST))[self.NODE]
        self.assertEqual((node['party'], node['party_source']), ('D', 'filing'))

    def test_each_cycle_reads_its_own_cycles_master_for_party(self):
        with tempfile.TemporaryDirectory() as directory:
            for folder, party in (('cn22', 'DEM'), ('cn24', 'REP')):
                (Path(directory) / folder).mkdir()
                (Path(directory) / folder / 'cn.txt').write_text(self.master('SMITH, ANN').replace('|REP|', f'|{party}|'),
                                                               encoding='ascii')
            saved, flow_data.CANDIDATE_MASTER_ROOT = flow_data.CANDIDATE_MASTER_ROOT, directory
            try:
                party = flow_data.master_party
                self.assertEqual(party(self.CAND, '2024', 'DEM'), ('R', 'master'))
                self.assertEqual(party(self.CAND, 2022, 'REP'), ('D', 'master'))
                self.assertEqual(party(self.CAND, '2023-2024', 'DEM'), ('R', 'master'), "Graph's cycle names its end year")
                self.assertEqual(party(self.CAND, '2026', 'GRE'), ('G', 'filing'),
                                 'no cn26: another cycle\'s master never answers; the filed party, flagged')
                self.assertEqual(party(self.CAND, '2026', ''), ('', ''))
                self.assertEqual(party('', '2024', 'DEM'), ('', ''))
            finally:
                flow_data.CANDIDATE_MASTER_ROOT = saved

    def test_graph_takes_party_from_the_masters_of_its_rows_cycles_and_flags_the_rest(self):
        """Regression (Graph): the party was the workbook's filed letter (D). cn26 says REP for the 2025-2026 rows."""
        republican = self.master('SMITH, ANN').replace('|2024|', '|2026|')
        for mode in ('direct', 'activity'):
            with self.subTest(mode=mode):
                body = self.graph(self.master('SMITH, ANN'), ['SMITH, ANN'], mode=mode, folders={'cn26': republican})
                row = next(row for row in body['candidates'] if row['id'] == self.CAND)
                self.assertEqual(row['party'], 'R')
                self.assertNotIn('party_source', row)
                self.assertNotIn('taken from the filing', body['candidate_note'])
                missing = self.graph(self.master('SMITH, ANN'), ['SMITH, ANN'], mode=mode)   # cn24 only
                row = next(row for row in missing['candidates'] if row['id'] == self.CAND)
                self.assertEqual((row['party'], row['party_source']), ('D', 'filing'))
                self.assertIn('Party for 1 candidate is taken from the filing', missing['candidate_note'])

    def test_a_graph_row_missing_from_the_master_keeps_the_workbooks_party_text(self):
        """The fallback is the filed party as written ("Other / Unclassified"), never a letter cut from it."""
        import graph_data
        with tempfile.TemporaryDirectory() as directory:
            saved, flow_data.CANDIDATE_MASTER_ROOT = flow_data.CANDIDATE_MASTER_ROOT, directory
            try:
                row = {'candidate_id': self.CAND, 'candidate': 'SMITH, ANN', 'party': 'Other / Unclassified',
                       'cycle': '2025-2026', 'amount': Decimal('1')}
                (named,), = graph_data.named_candidates([[row]])
            finally:
                flow_data.CANDIDATE_MASTER_ROOT = saved
        self.assertEqual((named['party'], named['party_source']), ('Other / Unclassified', 'filing'))

    def industry(self, master_lines):
        """industry_flow over a one-company workbook funding ROOT FUND (2024), whose saved snapshot is the base
        scenario with its first Schedule E row filed as DEM under the REP candidate H4XX01001."""
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            units = fixture_units()
            units[0]['expenditures'][0]['candidate_party'] = 'DEM'
            write_fixture(directory, units)
            path = Path(directory) / 'Acme.xlsx'
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = '2023-2024'
            for row in [['Acme giving'], ['Recipient facts'], [],
                        ['Fact Type', 'Company', 'Employee Tier', 'Amount', 'Recipient Category', 'Recipient Key',
                         'Include in Analysis', 'Contributor', 'Committee ID'],
                        ['Recipient', 'Acme', 'Corporate', 1500000, 'Super PAC', 'FEC:' + ROOT_ID, 'Yes', 'ACME CORP', ROOT_ID],
                        ['Master', 'Acme', 'Corporate', 1500000, 'Super PAC', None, 'Yes', None, None]]:
                sheet.append(row)
            maps = book.create_sheet('PAC & Candidate Map')
            for row in [['Mappings'], [], [], [], ['Cycle', 'Committee ID', 'Recipient Committee', 'Recipient Category',
                                                   'Direct Candidate / Sponsor', 'Candidate ID(s)', 'Direct Party', 'Relationship'],
                        ['2023-2024', ROOT_ID, 'ROOT FUND', 'Super PAC', '', '', '', 'No direct candidate/officeholder link']]:
                maps.append(row)
            activity = book.create_sheet('Downstream Activity')
            for row in [['Downstream'], [], ['1 source rows condensed to 1 committee-candidate-activity groups'], [],
                        ['Cycle', 'Source Committee ID', 'Candidate ID', 'Candidate / Officeholder', 'Candidate Party',
                         'Support / Oppose', 'Net Included Amount'],
                        ['2023-2024', ROOT_ID, self.CAND, 'SMITH, ANN', 'D', 'Support', 400000]]:
                activity.append(row)
            book.save(path)
            book.close()
            master = Path(directory) / 'master'
            (master / 'cn24').mkdir(parents=True)
            (master / 'cn24' / 'cn.txt').write_text(master_lines, encoding='ascii')
            saved = flow_data.CANDIDATE_MASTER_ROOT, downstream.ROOT
            flow_data.CANDIDATE_MASTER_ROOT, downstream.ROOT = master, Path(directory)
            try:
                built = flow_data.industry_flow('Acme', '2024', [path], top=8)
                body = flow_data.industry_payload('Acme', '2024', [path], top=8)
            finally:
                flow_data.CANDIDATE_MASTER_ROOT, downstream.ROOT = saved
        into = defaultdict(float)
        for link in built['graph'].links.values():
            if link['target'].startswith('outcome:'):
                into[link['target']] += float(link['value'])
        return dict(into), body

    def test_industry_flow_places_spending_by_the_masters_party(self):
        """Regression (industry flow): the first filed Schedule E row (DEM) put the REP candidate's $1,000,000 of
        support under Democrats. The master moves it to Republicans; the total is unchanged."""
        jones = 'S4YY00002|JONES, BOB|DEM|2024|YY|S|00|C|C|C90000011|1 MAIN ST||CITY|YY|20001\n'
        placed, body = self.industry(self.master('SMITH, ANN') + jones)
        self.assertEqual(placed, {'outcome:support:R': 1000000.0, 'outcome:oppose:D': 150000.0, 'outcome:contribution': 20000.0})
        self.assertFalse([note for note in body['notes'] if note['key'] == 'party_from_filing'])
        missing, body = self.industry(jones)
        self.assertEqual(missing['outcome:support:D'], 1000000.0, 'not in cn24: the first filing answers')
        self.assertEqual(sum(missing.values()), sum(placed.values()))
        self.assertEqual([note['text'] for note in body['notes'] if note['key'] == 'party_from_filing'],
                         [f"1 candidate's party taken from the filing: the {flow_data.cycle_label('2024')} FEC candidate master gives none"])

    def test_the_master_places_the_node_and_the_mixed_filing_is_disclosed(self):
        body = self.flow(self.master('SMITH, ANNABELLE ROSE'))
        node = by_id(body)[self.NODE]
        self.assertEqual((node['label'], node['state'], node['candidate_id']), ('Ann Smith', 'XX', self.CAND))
        self.assertEqual(node['name_as_filed'], 'SMITH, ANN', 'a filed spelling agrees, so the legal name is not shown')
        self.assertNotIn('name_master', node)
        self.assertEqual(node['cross_candidate_filings'], [self.note('k2')])
        self.assertEqual(self.received(body), 3500.0)
        self.assertTrue(body['reconciliation']['equal'])

    def note(self, sub_id='k2', named_id='', name='ZEDDY, ZOE', state='YY', amount=1000.0):
        return {'kind': 'cross_candidate_filing', 'sub_id': sub_id, 'amount': amount, 'date': '2023-10-02',
                'filer_committee_id': BRIDGE_ID, 'filed_candidate_id': self.CAND, 'filed_committee_id': 'C90000009',
                'named_recipient': name, 'named_office_state': state, 'named_recipient_committee': 'SMITH FOR CONGRESS',
                'named_candidate_id': named_id}

    def test_a_filing_that_names_another_candidate_is_a_tracked_note_on_the_node_the_link_and_the_flow(self):
        """Kept where FEC filed it: the note is listed and counted, and the node, link and reconciliation money
        do not move. The named candidate resolves to the one master candidate with that surname, office and state."""
        zeddy = 'H4YY05005|ZEDDY, ZOE|DEM|2024|YY|H|05|C|C|C90000010|1 MAIN ST||CITY|YY|20001\n'
        resolved = self.flow(self.master('SMITH, ANN') + zeddy)
        plain = self.flow(self.master('SMITH, ANN'))
        node = by_id(resolved)[self.NODE]
        self.assertEqual(node['cross_candidate_filings'], [self.note(named_id='H4YY05005')])
        into = [item for item in resolved['links'] if item['target'] == self.NODE]
        self.assertEqual([item['cross_candidate_filings'] for item in into], [[self.note(named_id='H4YY05005')]])
        self.assertEqual(into[0]['value'], 3500.0, 'the link still carries both filings')
        flow = resolved['cross_candidate_filings']
        self.assertEqual((flow['count'], flow['amount']), (1, 1000.0))
        self.assertEqual(flow['items'], [{**self.note(named_id='H4YY05005'), 'node': self.NODE, 'node_label': 'Ann Smith'}])
        for key in ('totals', 'reconciliation'):
            self.assertEqual(resolved[key], plain[key])
        self.assertEqual(self.received(resolved), 3500.0)
        self.assertEqual([(item['source'], item['target'], item['value']) for item in resolved['links']],
                         [(item['source'], item['target'], item['value']) for item in plain['links']])

    def test_the_named_candidate_is_blank_when_the_master_has_none_or_more_than_one_or_another_office(self):
        def named(extra):
            return by_id(self.flow(self.master('SMITH, ANN') + extra))[self.NODE]['cross_candidate_filings'][0]['named_candidate_id']
        line = 'H4YY0500%d|ZEDDY, ZOE|DEM|2024|YY|H|05|C|C|C90000010|1 MAIN ST||CITY|YY|20001\n'
        self.assertEqual(named(''), '', 'not in the master')
        self.assertEqual(named(line % 5 + line % 6), '', 'two candidates share the surname, office and state')
        self.assertEqual(named('S4YY00009|ZEDDY, ZOE|DEM|2024|YY|S|00|C|C|C90000010|1 MAIN ST||CITY|YY|20001\n'), '',
                         'a senator is not a House candidate')
        self.assertEqual(named(line.replace('YY|H', 'ZZ|H') % 5), '', 'another state')
        self.assertEqual(named(line % 5 + 'H4YY05006|OTHER, OLLIE|DEM|2024|YY|H|05|C|C|C90000010|1 MAIN ST||CITY|YY|20001\n'),
                         'H4YY05005', 'one match among other candidates')

    def test_a_node_with_no_other_name_has_no_note_and_the_flow_has_no_summary(self):
        body = self.flow(self.master('SMITH, ANN'), rows=(('k1', 'SMITH, ANN', 'XX', 2500.0),))
        self.assertNotIn('cross_candidate_filings', by_id(body)[self.NODE])
        self.assertTrue(all('cross_candidate_filings' not in item for item in body['links']))
        self.assertNotIn('cross_candidate_filings', body, 'absent when there is none, so the contract examples hold')

    def test_the_master_places_the_node_but_does_not_rename_a_filed_nickname(self):
        body = self.flow(self.master('SMITH, ANNABELLE ROSE', 'YY'), rows=(('k1', 'SMITH, ANN', 'XX', 2500.0),))
        node = by_id(body)[self.NODE]
        self.assertEqual((node['label'], node['state']), ('Ann Smith', 'YY'))
        self.assertNotIn('name_master', node)
        self.assertNotIn('cross_candidate_filings', node)

    def test_the_master_is_the_label_when_no_filed_spelling_agrees_with_it(self):
        body = self.flow(self.master('JONES, PAT'))
        node = by_id(body)[self.NODE]
        self.assertEqual((node['label'], node['state'], node['name_as_filed'], node['name_master']),
                         ('Pat Jones', 'XX', 'JONES, PAT', '2024'))
        self.assertEqual([item['named_recipient'] for item in node['cross_candidate_filings']], ['SMITH, ANN', 'ZEDDY, ZOE'])
        self.assertEqual(self.received(body), 3500.0)

    def test_the_master_names_a_node_whose_rows_carry_no_candidate_name_and_a_payee_name_is_no_other_filing(self):
        """A row with no candidate_name falls back to its payee ("Senate Eagle PAC" for Bill Hagerty): the master
        names it, and that payee name is never listed as filed under another candidate."""
        node = by_id(self.flow(self.master('SMITH, ANN'), rows=(('k1', None, 'XX', 2500.0), ('k2', None, 'XX', 1000.0))))[self.NODE]
        self.assertEqual((node['label'], node['name_master']), ('Ann Smith', '2024'))
        self.assertNotIn('cross_candidate_filings', node)
        mixed = by_id(self.flow(self.master('SMITH, ANN'), rows=(('k1', 'SMITH, ANN', 'XX', 2500.0), ('k2', None, 'XX', 1000.0))))
        self.assertNotIn('cross_candidate_filings', mixed[self.NODE])
        self.assertNotIn('name_master', mixed[self.NODE])

    def test_a_married_or_compound_filed_name_is_not_another_candidate(self):
        rows = (('k1', 'HINSON, ASHLEY', 'XX', 2500.0), ('k2', 'ARENHOLZ, ASHLEY HINSON', 'XX', 1000.0),
                ('k3', "O'BRIEN-HINSON, ASHLEY", 'XX', 500.0))
        node = by_id(self.flow(self.master('ARENHOLZ, ASHLEY HINSON'), rows=rows))[self.NODE]
        self.assertEqual(node['name_as_filed'], 'ARENHOLZ, ASHLEY HINSON')       # the master's own surname part wins
        self.assertNotIn('name_master', node)
        self.assertNotIn('cross_candidate_filings', node)
        only = by_id(self.flow(self.master('ARENHOLZ, ASHLEY HINSON'), rows=rows[:1]))[self.NODE]
        self.assertEqual((only['name_as_filed'], 'cross_candidate_filings' in only), ('HINSON, ASHLEY', False))

    def test_a_shared_given_name_is_not_agreement_with_the_master(self):
        """'JOHNSON, ANN' filed under SMITH, ANN's ID shares only the given name: the master labels the node and the
        filing is disclosed as naming another candidate."""
        body = self.flow(self.master('SMITH, ANN'), rows=(('k1', 'JOHNSON, ANN', 'XX', 2500.0),))
        node = by_id(body)[self.NODE]
        self.assertEqual((node['label'], node['name_as_filed'], node['name_master']), ('Ann Smith', 'SMITH, ANN', '2024'))
        self.assertEqual([item['named_recipient'] for item in node['cross_candidate_filings']], ['JOHNSON, ANN'])
        self.assertFalse(flow_data.same_person('SMITH, JAMES', 'JOHNSON, JAMES'))
        self.assertTrue(flow_data.same_person('HINSON, ASHLEY', 'ARENHOLZ, ASHLEY HINSON'))

    def test_without_a_master_the_rows_filed_in_the_ids_own_state_name_the_node(self):
        """Same name length and one vote each: today's alphabetical tie-break gives ZEDDY here."""
        reference = self.flow(self.master('SMITH, ANN'))
        body = self.flow()
        node = by_id(body)[self.NODE]
        self.assertEqual((node['label'], node['state']), ('Ann Smith', 'XX'))
        self.assertNotIn('name_master', node)
        self.assertEqual([item['named_recipient'] for item in node['cross_candidate_filings']], ['ZEDDY, ZOE'])
        self.assertEqual(self.received(body), 3500.0)
        self.assertEqual(body['reconciliation'], reference['reconciliation'])

    def test_an_id_state_that_no_filed_row_shares_gives_way_to_the_filed_state(self):
        body = self.flow(rows=(('k1', 'SMITH, ANN', 'YY', 2500.0), ('k2', 'ZEDDY, ZOE', 'ZZ', 1000.0)))
        node = by_id(body)[self.NODE]
        self.assertEqual((node['label'], node['state']), ('Ann Smith', 'YY'))       # the larger amount wins the tie
        self.assertEqual(self.received(body), 3500.0)

    def test_the_candidate_master_of_the_flows_cycle_is_read_first(self):
        with tempfile.TemporaryDirectory() as directory:
            for folder, name, state in (('cn22', 'OLD, OLLIE', 'AA'), ('cn24', 'NEW, NORA', 'BB'), ('cn26', 'NEWEST, NED', 'CC')):
                (Path(directory) / folder).mkdir()
                (Path(directory) / folder / 'cn.txt').write_text(f'H4XX01001|{name}|REP|2024|{state}|H|01|I|C|C1|\n', encoding='ascii')
            saved, flow_data.CANDIDATE_MASTER_ROOT = flow_data.CANDIDATE_MASTER_ROOT, directory
            try:
                for cycle, year in ((2022, '2022'), ('2024', '2024'), (2028, '2026'), (None, '2026')):
                    self.assertEqual([item[0] for item in flow_data.candidate_master(cycle)][0], year, cycle)
                self.assertEqual(flow_data.candidate_master(2024)[0][1]['H4XX01001'], ('NEW, NORA', 'BB', 'REP'))
            finally:
                flow_data.CANDIDATE_MASTER_ROOT = saved

    def test_name_words_ignore_case_accents_apostrophes_and_short_words(self):
        words = flow_data.name_words("O'Brien, JOSÉ de la Cruz Jr")
        self.assertTrue({'OBRIEN', 'JOSE', 'CRUZ'} <= words and not words & {'DE', 'LA', 'JR'}, words)
        self.assertTrue(flow_data.name_words('VAN EPPS, MATT') & flow_data.name_words('VANEPPS, MATTHEW ROBERT'))
        self.assertEqual(flow_data.name_words(None), set())


if __name__ == '__main__':
    unittest.main()
